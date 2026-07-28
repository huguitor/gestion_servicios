from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError
from django.db import transaction
from django.test import TestCase
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APITestCase

from clientes.models import Cliente
from comprobantes.models import Comprobante
from presupuestos.models import Presupuesto
from remitos.models import Remito

from cobranzas.models import Cobro, FacturaCobranza, SeguimientoCobranza
from cobranzas.services import registrar_cobro
from cobranzas.views import FacturaCobranzaViewSet


class CobranzasFixturesMixin:
    def crear_cliente(self, nombre="Cliente", plazo=0):
        return Cliente.objects.create(
            tipo="juridica",
            nombre=nombre,
            plazo_cobro_dias=plazo,
        )

    def crear_factura(self, cliente=None, **overrides):
        datos = {
            "cliente": cliente or self.cliente,
            "fecha_factura": date(2026, 7, 1),
            "tipo_comprobante": FacturaCobranza.TIPO_FACTURA_A,
            "punto_venta": 1,
            "numero_factura": 150,
            "total": Decimal("1000.00"),
            "creado_por": self.usuario,
        }
        datos.update(overrides)
        return FacturaCobranza.objects.create(**datos)


class FacturaCobranzaModelTests(CobranzasFixturesMixin, TestCase):
    def setUp(self):
        self.usuario = get_user_model().objects.create_user("cobranzas-model")
        self.cliente = self.crear_cliente(plazo=15)

    def test_copia_plazo_y_calcula_fecha_desde_envio_en_api_de_dominio(self):
        factura = self.crear_factura(
            plazo_cobro_dias=self.cliente.plazo_cobro_dias,
            fecha_envio=date(2026, 7, 5),
        )
        self.assertEqual(factura.plazo_cobro_dias, 15)
        self.assertEqual(factura.fecha_estimada_cobro, date(2026, 7, 20))

        self.cliente.plazo_cobro_dias = 60
        self.cliente.save()
        factura.refresh_from_db()
        self.assertEqual(factura.plazo_cobro_dias, 15)

    def test_creacion_fuera_de_api_copia_plazo_y_calcula_fecha(self):
        factura = self.crear_factura(
            fecha_envio=date(2026, 7, 5),
        )
        self.assertEqual(factura.plazo_cobro_dias, 15)
        self.assertEqual(factura.fecha_estimada_cobro, date(2026, 7, 20))

    def test_fecha_estimada_usa_fecha_factura_como_fallback(self):
        factura = self.crear_factura(plazo_cobro_dias=10)
        self.assertEqual(factura.fecha_estimada_cobro, date(2026, 7, 11))

    def test_fecha_manual_no_se_sobrescribe(self):
        factura = self.crear_factura(
            plazo_cobro_dias=10,
            fecha_estimada_cobro=date(2026, 9, 30),
            fecha_estimada_manual=True,
        )
        factura.observaciones = "Edición sin relación con el vencimiento"
        factura.plazo_cobro_dias = 45
        factura.save()
        self.assertEqual(factura.fecha_estimada_cobro, date(2026, 9, 30))

    def test_total_y_plazo_tienen_restricciones(self):
        with self.assertRaises(ValidationError):
            self.crear_factura(total=Decimal("0.00"))

    def test_identificacion_de_factura_es_unica(self):
        self.crear_factura()
        with self.assertRaises(ValidationError):
            self.crear_factura()

    def test_presupuesto_debe_pertenecer_al_cliente(self):
        otro_cliente = self.crear_cliente("Otro")
        comprobante = Comprobante.objects.create(
            tipo="PRES",
            serie="00001",
            numero_inicial=1,
            numero_final=999999,
            proximo_numero=1,
        )
        presupuesto = Presupuesto.objects.create(
            cliente=otro_cliente,
            comprobante=comprobante,
            creado_por=self.usuario,
        )
        with self.assertRaises(ValidationError):
            self.crear_factura(presupuesto=presupuesto)

    def test_cambiar_cliente_revalida_presupuesto_y_remitos_existentes(self):
        otro_cliente = self.crear_cliente("Otro")
        comp_presupuesto = Comprobante.objects.create(
            tipo="PRES",
            serie="00001",
            numero_inicial=1,
            numero_final=999999,
            proximo_numero=1,
        )
        presupuesto = Presupuesto.objects.create(
            cliente=self.cliente,
            comprobante=comp_presupuesto,
            creado_por=self.usuario,
        )
        comp_remito = Comprobante.objects.create(
            tipo="REMI",
            serie="00001",
            numero_inicial=1,
            numero_final=999999,
            proximo_numero=1,
        )
        remito = Remito.objects.create(
            cliente=self.cliente,
            comprobante=comp_remito,
            creado_por=self.usuario,
        )
        factura = self.crear_factura(presupuesto=presupuesto)
        factura.remitos.add(remito)

        factura.cliente = otro_cliente
        with self.assertRaises(ValidationError) as contexto:
            factura.save()
        self.assertIn("presupuesto", contexto.exception.message_dict)
        self.assertIn("cliente", contexto.exception.message_dict)

    def test_remito_ajeno_se_rechaza_tambien_desde_gestor_m2m(self):
        otro_cliente = self.crear_cliente("Otro")
        comprobante = Comprobante.objects.create(
            tipo="REMI",
            serie="00001",
            numero_inicial=1,
            numero_final=999999,
            proximo_numero=1,
        )
        remito = Remito.objects.create(
            cliente=otro_cliente,
            comprobante=comprobante,
            creado_por=self.usuario,
        )
        factura = self.crear_factura()
        with self.assertRaises(ValidationError), transaction.atomic():
            factura.remitos.add(remito)
        with self.assertRaises(ValidationError), transaction.atomic():
            remito.facturas_cobranza.add(factura)

    def test_recalculo_se_incluye_al_guardar_con_update_fields(self):
        factura = self.crear_factura(plazo_cobro_dias=10)
        factura.fecha_envio = date(2026, 7, 5)
        factura.save(update_fields=["fecha_envio"])
        factura.refresh_from_db()
        self.assertEqual(factura.fecha_estimada_cobro, date(2026, 7, 15))

    def test_estados_y_saldos_son_calculados(self):
        factura = self.crear_factura()
        self.assertEqual(factura.estado, "pendiente")
        registrar_cobro(
            factura=factura,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 7, 10),
            importe=Decimal("400.00"),
            medio_pago="transferencia",
        )
        self.assertEqual(factura.total_cobrado, Decimal("400.00"))
        self.assertEqual(factura.saldo_pendiente, Decimal("600.00"))
        self.assertEqual(factura.estado, "parcial")
        registrar_cobro(
            factura=factura,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 7, 11),
            importe=Decimal("600.00"),
            medio_pago="cheque",
        )
        self.assertEqual(factura.estado, "pagado")
        self.assertEqual(factura.fecha_ultimo_cobro, date(2026, 7, 11))

    def test_identidad_queda_bloqueada_pero_total_se_puede_editar_despues_de_cobrar(self):
        factura = self.crear_factura()
        registrar_cobro(
            factura=factura,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 7, 10),
            importe=Decimal("400.00"),
            medio_pago="transferencia",
        )

        cambios_bloqueados = {
            "cliente": self.crear_cliente("Otro"),
            "fecha_factura": date(2026, 7, 2),
            "tipo_comprobante": FacturaCobranza.TIPO_FACTURA_B,
            "punto_venta": 2,
            "numero_factura": 999,
        }
        for campo, valor in cambios_bloqueados.items():
            with self.subTest(campo=campo):
                factura.refresh_from_db()
                setattr(factura, campo, valor)
                with self.assertRaises(ValidationError):
                    factura.save()

        factura.refresh_from_db()
        factura.total = Decimal("1200.02")
        factura.save()
        factura.refresh_from_db()
        self.assertEqual(factura.total, Decimal("1200.02"))
        factura.fecha_estimada_manual = True
        factura.fecha_estimada_cobro = date(2026, 8, 31)
        factura.save()
        self.assertEqual(factura.fecha_estimada_cobro, date(2026, 8, 31))

    def test_no_permite_reducir_total_ni_cambiar_remitos_despues_de_cobrar(self):
        comprobante = Comprobante.objects.create(
            tipo="REMI",
            serie="00001",
            numero_inicial=1,
            numero_final=999999,
            proximo_numero=1,
        )
        remito = Remito.objects.create(
            cliente=self.cliente,
            comprobante=comprobante,
            creado_por=self.usuario,
        )
        factura = self.crear_factura()
        factura.remitos.add(remito)
        registrar_cobro(
            factura=factura,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 7, 10),
            importe=Decimal("400.00"),
            medio_pago="transferencia",
        )
        factura.total = Decimal("399.00")
        with self.assertRaises(ValidationError):
            factura.save()
        with self.assertRaises(ValidationError):
            factura.remitos.clear()

    def test_anotaciones_resuelven_total_y_ultimo_cobro_sin_consultas_extra(self):
        factura = self.crear_factura()
        registrar_cobro(
            factura=factura,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 7, 10),
            importe=Decimal("250.00"),
            medio_pago="transferencia",
        )
        view = FacturaCobranzaViewSet()
        view.action = "list"
        factura_anotada = view.get_queryset().get(pk=factura.pk)
        with self.assertNumQueries(0):
            self.assertEqual(factura_anotada.total_cobrado, Decimal("250.00"))
            self.assertEqual(factura_anotada.fecha_ultimo_cobro, date(2026, 7, 10))

    @patch("cobranzas.models.timezone.localdate", return_value=date(2026, 7, 20))
    def test_semaforo_y_vencimiento_incluyen_limites(self, _localdate):
        pagada = self.crear_factura(
            numero_factura=1,
            fecha_estimada_cobro=date(2026, 7, 10),
            fecha_estimada_manual=True,
            total=Decimal("10.00"),
        )
        registrar_cobro(
            factura=pagada,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 7, 10),
            importe=Decimal("10.00"),
            medio_pago="efectivo",
        )
        self.assertEqual(pagada.semaforo, "gris")

        verde = self.crear_factura(
            numero_factura=2,
            fecha_estimada_cobro=date(2026, 7, 28),
            fecha_estimada_manual=True,
        )
        amarillo_siete = self.crear_factura(
            numero_factura=3,
            fecha_estimada_cobro=date(2026, 7, 27),
            fecha_estimada_manual=True,
        )
        amarillo_hoy = self.crear_factura(
            numero_factura=4,
            fecha_estimada_cobro=date(2026, 7, 20),
            fecha_estimada_manual=True,
        )
        roja = self.crear_factura(
            numero_factura=5,
            fecha_estimada_cobro=date(2026, 7, 19),
            fecha_estimada_manual=True,
        )
        registrar_cobro(
            factura=roja,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 7, 15),
            importe=Decimal("100.00"),
            medio_pago="transferencia",
        )

        self.assertEqual(verde.semaforo, "verde")
        self.assertEqual(amarillo_siete.semaforo, "amarillo")
        self.assertEqual(amarillo_hoy.semaforo, "amarillo")
        self.assertEqual(roja.estado, "parcial")
        self.assertTrue(roja.vencida)
        self.assertEqual(roja.semaforo, "rojo")


class CobroServiceTests(CobranzasFixturesMixin, TestCase):
    def setUp(self):
        self.usuario = get_user_model().objects.create_user("cobranzas-cobro")
        self.cliente = self.crear_cliente()
        self.factura = self.crear_factura()

    def test_rechaza_importes_no_positivos(self):
        for importe in (Decimal("0.00"), Decimal("-1.00")):
            with self.subTest(importe=importe), self.assertRaises(ValidationError):
                registrar_cobro(
                    factura=self.factura,
                    registrado_por=self.usuario,
                    fecha_cobro=date(2026, 7, 10),
                    importe=importe,
                    medio_pago="transferencia",
                )
        self.assertFalse(Cobro.objects.exists())

    def test_rechaza_sobrepago_y_revierte_la_operacion(self):
        registrar_cobro(
            factura=self.factura,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 7, 10),
            importe=Decimal("900.00"),
            medio_pago="transferencia",
        )
        with self.assertRaises(ValidationError):
            registrar_cobro(
                factura=self.factura,
                registrado_por=self.usuario,
                fecha_cobro=date(2026, 7, 11),
                importe=Decimal("101.00"),
                medio_pago="transferencia",
            )
        self.assertEqual(Cobro.objects.count(), 1)
        self.assertEqual(self.factura.total_cobrado, Decimal("900.00"))


class CobranzasAPITests(CobranzasFixturesMixin, APITestCase):
    def setUp(self):
        self.usuario = get_user_model().objects.create_user(
            "cobranzas-api",
            is_staff=True,
        )
        permisos = Permission.objects.filter(content_type__app_label="cobranzas")
        self.usuario.user_permissions.add(*permisos)
        self.client.force_authenticate(self.usuario)
        self.cliente = self.crear_cliente(plazo=20)

    def payload_factura(self, **overrides):
        payload = {
            "cliente": self.cliente.id,
            "fecha_factura": "2026-07-01",
            "tipo_comprobante": "factura_a",
            "punto_venta": 1,
            "numero_factura": 150,
            "total": "1000.00",
        }
        payload.update(overrides)
        return payload

    def test_creacion_copia_plazo_y_devuelve_calculos(self):
        response = self.client.post(
            reverse("cobranzas-facturas-list"),
            self.payload_factura(fecha_envio="2026-07-05"),
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(response.data["plazo_cobro_dias"], 20)
        self.assertEqual(response.data["fecha_estimada_cobro"], "2026-07-25")
        self.assertEqual(response.data["numero_completo"], "00001-00000150")
        self.assertEqual(response.data["estado"], "pendiente")

    def test_importes_se_conservan_exactamente_en_todo_el_flujo_api(self):
        valores = ("0.01", "0.02", "10.10", "100.99", "125430.52", "123456.78")
        for indice, valor in enumerate(valores, start=1):
            with self.subTest(valor=valor):
                response = self.client.post(
                    reverse("cobranzas-facturas-list"),
                    self.payload_factura(numero_factura=200 + indice, total=valor),
                    format="json",
                )
                self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
                self.assertEqual(response.data["total"], valor)
                factura = FacturaCobranza.objects.get(pk=response.data["id"])
                self.assertIsInstance(factura.total, Decimal)
                self.assertEqual(factura.total, Decimal(valor))
                recuperada = self.client.get(
                    reverse("cobranzas-facturas-detail", args=[factura.id])
                )
                self.assertEqual(recuperada.data["total"], valor)

    def test_actualizar_importe_conserva_centavos_y_permiso_existente(self):
        factura = self.crear_factura(total=Decimal("125430.52"))
        response = self.client.patch(
            reverse("cobranzas-facturas-detail", args=[factura.id]),
            {"total": "125430.50"},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data["total"], "125430.50")
        self.assertEqual(response.data["saldo_pendiente"], "125430.50")
        factura.refresh_from_db()
        self.assertEqual(factura.total, Decimal("125430.50"))

    def test_nota_credito_aplica_signo_sin_contarse_como_cobro(self):
        factura = self.crear_factura(total=Decimal("100000.00"))
        registrar_cobro(
            factura=factura,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 7, 10),
            importe=Decimal("80000.00"),
            medio_pago="transferencia",
        )
        response = self.client.post(
            reverse("cobranzas-facturas-list"),
            self.payload_factura(
                numero_factura=901,
                tipo_comprobante="nota_credito",
                total="20000.00",
                comprobante_original=factura.id,
            ),
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(response.data["total"], "20000.00")
        self.assertEqual(response.data["importe_con_efecto"], "-20000.00")
        self.assertEqual(response.data["total_cobrado"], "0.00")
        self.assertEqual(response.data["estado"], "aplicada")
        factura.refresh_from_db()
        self.assertEqual(factura.total_notas_credito, Decimal("20000.00"))
        self.assertEqual(factura.total_cobrado, Decimal("80000.00"))
        self.assertEqual(factura.saldo_pendiente, Decimal("0.00"))
        self.assertEqual(factura.estado, "pagado")

    def test_nota_credito_rechaza_importes_no_positivos_y_exceso(self):
        factura = self.crear_factura(total=Decimal("100.00"))
        for indice, valor in enumerate(("0.00", "-1.00"), start=1):
            response = self.client.post(
                reverse("cobranzas-facturas-list"),
                self.payload_factura(
                    numero_factura=920 + indice,
                    tipo_comprobante="nota_credito",
                    total=valor,
                    comprobante_original=factura.id,
                ),
                format="json",
            )
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        exceso = self.client.post(
            reverse("cobranzas-facturas-list"),
            self.payload_factura(
                numero_factura=930,
                tipo_comprobante="nota_credito",
                total="100.01",
                comprobante_original=factura.id,
            ),
            format="json",
        )
        self.assertEqual(exceso.status_code, status.HTTP_400_BAD_REQUEST, exceso.data)

    def test_caso_integrado_centavos_nota_credito_y_cobro(self):
        factura = self.crear_factura(total=Decimal("125430.52"))
        nota = self.client.post(
            reverse("cobranzas-facturas-list"),
            self.payload_factura(
                numero_factura=940,
                tipo_comprobante="nota_credito",
                total="0.02",
                comprobante_original=factura.id,
            ),
            format="json",
        )
        self.assertEqual(nota.status_code, status.HTTP_201_CREATED, nota.data)
        cobro = self.client.post(
            reverse("cobranzas-facturas-cobros", args=[factura.id]),
            {
                "fecha_cobro": "2026-07-10",
                "importe": "125430.50",
                "medio_pago": "transferencia",
            },
            format="json",
        )
        self.assertEqual(cobro.status_code, status.HTTP_201_CREATED, cobro.data)
        detalle = self.client.get(
            reverse("cobranzas-facturas-detail", args=[factura.id])
        )
        self.assertEqual(detalle.data["total"], "125430.52")
        self.assertEqual(detalle.data["total_notas_credito"], "0.02")
        self.assertEqual(detalle.data["total_cobrado"], "125430.50")
        self.assertEqual(detalle.data["saldo_pendiente"], "0.00")

    def test_pago_parcial_conserva_saldo_de_dos_centavos(self):
        factura = self.crear_factura(total=Decimal("125430.52"))
        cobro = self.client.post(
            reverse("cobranzas-facturas-cobros", args=[factura.id]),
            {
                "fecha_cobro": "2026-07-10",
                "importe": "125430.50",
                "medio_pago": "transferencia",
            },
            format="json",
        )
        self.assertEqual(cobro.status_code, status.HTTP_201_CREATED, cobro.data)
        detalle = self.client.get(
            reverse("cobranzas-facturas-detail", args=[factura.id])
        )
        self.assertEqual(detalle.data["saldo_pendiente"], "0.02")
        self.assertEqual(detalle.data["estado"], "parcial")

    def test_fecha_enviada_manualmente_activa_marca(self):
        response = self.client.post(
            reverse("cobranzas-facturas-list"),
            self.payload_factura(fecha_estimada_cobro="2026-09-30"),
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertTrue(response.data["fecha_estimada_manual"])
        self.assertEqual(response.data["fecha_estimada_cobro"], "2026-09-30")

    def test_valida_presupuesto_y_varios_remitos_del_cliente(self):
        otro = self.crear_cliente("Cliente ajeno")
        comp_pres = Comprobante.objects.create(
            tipo="PRES",
            serie="00001",
            numero_inicial=1,
            numero_final=999999,
            proximo_numero=1,
        )
        presupuesto = Presupuesto.objects.create(
            cliente=self.cliente,
            comprobante=comp_pres,
            creado_por=self.usuario,
        )
        presupuesto_ajeno = Presupuesto.objects.create(
            cliente=otro,
            comprobante=comp_pres,
            creado_por=self.usuario,
        )
        comp_remito = Comprobante.objects.create(
            tipo="REMI",
            serie="00001",
            numero_inicial=1,
            numero_final=999999,
            proximo_numero=1,
        )
        remitos = [
            Remito.objects.create(
                cliente=self.cliente,
                comprobante=comp_remito,
                creado_por=self.usuario,
            )
            for _ in range(2)
        ]
        remito_ajeno = Remito.objects.create(
            cliente=otro,
            comprobante=comp_remito,
            creado_por=self.usuario,
        )

        correcto = self.client.post(
            reverse("cobranzas-facturas-list"),
            self.payload_factura(
                presupuesto=presupuesto.id,
                remitos=[remito.id for remito in remitos],
            ),
            format="json",
        )
        self.assertEqual(correcto.status_code, status.HTTP_201_CREATED, correcto.data)
        self.assertEqual(len(correcto.data["remitos"]), 2)
        self.assertEqual(correcto.data["presupuesto_referencia"], "00001-000001")

        presupuesto_invalido = self.client.post(
            reverse("cobranzas-facturas-list"),
            self.payload_factura(
                numero_factura=151,
                presupuesto=presupuesto_ajeno.id,
            ),
            format="json",
        )
        self.assertEqual(presupuesto_invalido.status_code, status.HTTP_400_BAD_REQUEST)

        remito_invalido = self.client.post(
            reverse("cobranzas-facturas-list"),
            self.payload_factura(numero_factura=152, remitos=[remito_ajeno.id]),
            format="json",
        )
        self.assertEqual(remito_invalido.status_code, status.HTTP_400_BAD_REQUEST)

    def test_registro_de_cobro_y_seguimiento(self):
        factura = self.crear_factura()
        cobro = self.client.post(
            reverse("cobranzas-facturas-cobros", args=[factura.id]),
            {
                "fecha_cobro": "2026-07-10",
                "importe": "400.00",
                "medio_pago": "transferencia",
            },
            format="json",
        )
        self.assertEqual(cobro.status_code, status.HTTP_201_CREATED, cobro.data)
        detalle = self.client.get(
            reverse("cobranzas-facturas-detail", args=[factura.id])
        )
        self.assertEqual(detalle.data["total_cobrado"], "400.00")
        self.assertEqual(detalle.data["saldo_pendiente"], "600.00")
        self.assertEqual(detalle.data["estado"], "parcial")

        seguimiento = self.client.post(
            reverse("cobranzas-facturas-seguimientos", args=[factura.id]),
            {
                "fecha": "2026-07-11",
                "tipo": "reclamo",
                "detalle": "Se reclamó el pago.",
            },
            format="json",
        )
        self.assertEqual(
            seguimiento.status_code, status.HTTP_201_CREATED, seguimiento.data
        )
        self.assertEqual(SeguimientoCobranza.objects.count(), 1)

    def test_listado_no_carga_historiales_completos(self):
        factura = self.crear_factura()
        registrar_cobro(
            factura=factura,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 7, 10),
            importe=Decimal("100.00"),
            medio_pago="transferencia",
        )
        response = self.client.get(reverse("cobranzas-facturas-list"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertNotIn("cobros", response.data[0])
        self.assertNotIn("seguimientos", response.data[0])

    def test_filtros_fecha_son_inclusivos_y_validan_el_rango(self):
        self.crear_factura(numero_factura=301, fecha_factura=date(2026, 6, 30))
        self.crear_factura(numero_factura=302, fecha_factura=date(2026, 7, 1))
        self.crear_factura(numero_factura=303, fecha_factura=date(2026, 7, 31))
        self.crear_factura(numero_factura=304, fecha_factura=date(2026, 8, 1))
        url = reverse("cobranzas-facturas-list")

        desde = self.client.get(url, {"fecha_desde": "2026-07-31"})
        self.assertEqual(
            {row["numero_factura"] for row in desde.data},
            {303, 304},
        )
        hasta = self.client.get(url, {"fecha_hasta": "2026-07-01"})
        self.assertEqual(
            {row["numero_factura"] for row in hasta.data},
            {301, 302},
        )
        rango = self.client.get(
            url,
            {"fecha_desde": "2026-07-01", "fecha_hasta": "2026-07-31"},
        )
        self.assertEqual(
            {row["numero_factura"] for row in rango.data},
            {302, 303},
        )
        invalido = self.client.get(
            url,
            {"fecha_desde": "2026-08-01", "fecha_hasta": "2026-07-01"},
        )
        self.assertEqual(invalido.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("fecha_hasta", invalido.data)

    def test_fechas_se_combinan_con_cliente_estado_y_busqueda(self):
        otro = self.crear_cliente("Buscado Especial")
        self.crear_factura(
            numero_factura=310,
            fecha_factura=date(2026, 7, 10),
            cliente=otro,
        )
        parcial = self.crear_factura(
            numero_factura=311,
            fecha_factura=date(2026, 7, 11),
            cliente=otro,
        )
        registrar_cobro(
            factura=parcial,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 7, 12),
            importe=Decimal("10.10"),
            medio_pago="transferencia",
        )
        self.crear_factura(numero_factura=312, fecha_factura=date(2026, 8, 1))
        params = {
            "fecha_desde": "2026-07-01",
            "fecha_hasta": "2026-07-31",
            "cliente": otro.id,
            "estado": "parcial",
            "search": "Especial",
        }
        response = self.client.get(reverse("cobranzas-facturas-list"), params)
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual([row["numero_factura"] for row in response.data], [311])

    def test_permisos_especificos_de_cobros_y_seguimientos(self):
        factura = self.crear_factura()
        usuario = get_user_model().objects.create_user(
            "permisos-cobranzas",
            is_staff=True,
        )
        usuario.user_permissions.add(
            Permission.objects.get(
                content_type__app_label="cobranzas",
                codename="view_facturacobranza",
            )
        )
        self.client.force_authenticate(usuario)

        cobros_url = reverse("cobranzas-facturas-cobros", args=[factura.id])
        seguimientos_url = reverse(
            "cobranzas-facturas-seguimientos", args=[factura.id]
        )
        self.assertEqual(self.client.get(cobros_url).status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(
            self.client.post(
                cobros_url,
                {
                    "fecha_cobro": "2026-07-10",
                    "importe": "100.00",
                    "medio_pago": "transferencia",
                },
                format="json",
            ).status_code,
            status.HTTP_403_FORBIDDEN,
        )
        self.assertEqual(
            self.client.get(seguimientos_url).status_code,
            status.HTTP_403_FORBIDDEN,
        )
        self.assertEqual(
            self.client.post(
                seguimientos_url,
                {"fecha": "2026-07-10", "tipo": "nota", "detalle": "Prueba"},
                format="json",
            ).status_code,
            status.HTTP_403_FORBIDDEN,
        )

        permisos = Permission.objects.filter(
            content_type__app_label="cobranzas",
            codename__in=[
                "view_cobro",
                "registrar_cobro",
                "view_seguimientocobranza",
                "agregar_seguimiento",
            ],
        )
        usuario.user_permissions.add(*permisos)
        usuario = get_user_model().objects.get(pk=usuario.pk)
        self.client.force_authenticate(usuario)
        self.assertEqual(self.client.get(cobros_url).status_code, status.HTTP_200_OK)
        self.assertEqual(
            self.client.get(seguimientos_url).status_code,
            status.HTTP_200_OK,
        )

    def test_delete_no_esta_habilitado(self):
        factura = self.crear_factura()
        response = self.client.delete(
            reverse("cobranzas-facturas-detail", args=[factura.id])
        )
        self.assertEqual(response.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)

    def test_autenticacion_y_permisos(self):
        self.client.force_authenticate(user=None)
        self.assertEqual(
            self.client.get(reverse("cobranzas-facturas-list")).status_code,
            status.HTTP_401_UNAUTHORIZED,
        )
        sin_permiso = get_user_model().objects.create_user(
            "sin-permiso-cobranzas",
            is_staff=True,
        )
        self.client.force_authenticate(sin_permiso)
        self.assertEqual(
            self.client.get(reverse("cobranzas-facturas-list")).status_code,
            status.HTTP_403_FORBIDDEN,
        )

    @patch("cobranzas.models.timezone.localdate", return_value=date(2026, 7, 20))
    def test_filtros_estado_semaforo_y_vencidas(self, _localdate):
        self.crear_factura(
            numero_factura=1,
            fecha_estimada_cobro=date(2026, 7, 19),
            fecha_estimada_manual=True,
        )
        parcial = self.crear_factura(
            numero_factura=2,
            fecha_estimada_cobro=date(2026, 7, 27),
            fecha_estimada_manual=True,
        )
        registrar_cobro(
            factura=parcial,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 7, 15),
            importe=Decimal("100.00"),
            medio_pago="transferencia",
        )
        self.crear_factura(
            numero_factura=3,
            fecha_estimada_cobro=date(2026, 7, 29),
            fecha_estimada_manual=True,
        )

        parcial_response = self.client.get(
            reverse("cobranzas-facturas-list"), {"estado": "parcial"}
        )
        self.assertEqual(len(parcial_response.data), 1)
        rojas = self.client.get(
            reverse("cobranzas-facturas-list"), {"semaforo": "rojo"}
        )
        self.assertEqual(len(rojas.data), 1)
        vencidas = self.client.get(
            reverse("cobranzas-facturas-list"), {"vencidas": "true"}
        )
        self.assertEqual(len(vencidas.data), 1)

    @patch("cobranzas.views.timezone.localdate", return_value=date(2026, 7, 20))
    def test_dashboard_devuelve_solo_indicadores_acordados(self, _localdate):
        vencida = self.crear_factura(
            numero_factura=1,
            fecha_estimada_cobro=date(2026, 7, 19),
            fecha_estimada_manual=True,
        )
        proxima = self.crear_factura(
            numero_factura=2,
            fecha_estimada_cobro=date(2026, 7, 25),
            fecha_estimada_manual=True,
        )
        registrar_cobro(
            factura=proxima,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 7, 10),
            importe=Decimal("250.00"),
            medio_pago="transferencia",
        )
        pagada = self.crear_factura(numero_factura=3, total=Decimal("100.00"))
        registrar_cobro(
            factura=pagada,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 7, 11),
            importe=Decimal("100.00"),
            medio_pago="efectivo",
        )

        response = self.client.get(reverse("cobranzas-facturas-dashboard"))
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(Decimal(response.data["total_pendiente"]), Decimal("1750.00"))
        self.assertEqual(Decimal(response.data["total_vencido"]), Decimal("1000.00"))
        self.assertEqual(Decimal(response.data["total_cobrado_mes"]), Decimal("350.00"))
        self.assertEqual(
            response.data["cantidades_por_estado"],
            {"pendiente": 1, "parcial": 1, "pagado": 1},
        )
        self.assertEqual(len(response.data["proximas_a_vencer"]), 1)

    @patch("cobranzas.views.timezone.localdate", return_value=date(2026, 8, 20))
    def test_dashboard_respeta_rango_y_cobros_de_comprobantes_filtrados(
        self, _localdate
    ):
        factura = self.crear_factura(
            numero_factura=401,
            fecha_factura=date(2026, 7, 1),
            total=Decimal("125430.52"),
            fecha_estimada_cobro=date(2026, 7, 10),
            fecha_estimada_manual=True,
        )
        registrar_cobro(
            factura=factura,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 8, 5),
            importe=Decimal("10.10"),
            medio_pago="transferencia",
        )
        self.crear_factura(
            numero_factura=402,
            fecha_factura=date(2026, 7, 15),
            tipo_comprobante=FacturaCobranza.TIPO_NOTA_CREDITO,
            total=Decimal("0.02"),
        )
        self.crear_factura(
            numero_factura=403,
            fecha_factura=date(2026, 8, 1),
            total=Decimal("123456.78"),
        )
        response = self.client.get(
            reverse("cobranzas-facturas-dashboard"),
            {"fecha_desde": "2026-07-01", "fecha_hasta": "2026-07-31"},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(Decimal(response.data["total_facturado"]), Decimal("125430.52"))
        self.assertEqual(Decimal(response.data["total_notas_credito"]), Decimal("0.02"))
        self.assertEqual(Decimal(response.data["total_cobrado_mes"]), Decimal("10.10"))
        self.assertEqual(Decimal(response.data["total_pendiente"]), Decimal("125420.40"))
        self.assertEqual(Decimal(response.data["total_vencido"]), Decimal("125420.42"))
        self.assertEqual(
            response.data["cantidades_por_estado"],
            {"pendiente": 0, "parcial": 1, "pagado": 0},
        )

    def test_exportacion_excel_respeta_filtros_cabeceras_y_centavos(self):
        otro = self.crear_cliente("Cliente Excel")
        incluida = self.crear_factura(
            cliente=otro,
            numero_factura=501,
            fecha_factura=date(2026, 7, 31),
            total=Decimal("125430.52"),
            orden_compra="OC-XLSX",
        )
        registrar_cobro(
            factura=incluida,
            registrado_por=self.usuario,
            fecha_cobro=date(2026, 8, 1),
            importe=Decimal("0.02"),
            medio_pago="transferencia",
        )
        self.crear_factura(
            cliente=otro,
            numero_factura=503,
            fecha_factura=date(2026, 7, 31),
            tipo_comprobante=FacturaCobranza.TIPO_NOTA_CREDITO,
            total=Decimal("0.01"),
            comprobante_original=incluida,
            orden_compra="OC-XLSX",
        )
        self.crear_factura(
            numero_factura=502,
            fecha_factura=date(2026, 8, 1),
            total=Decimal("123456.78"),
        )
        url = reverse("cobranzas-facturas-exportar-excel")
        response = self.client.get(
            url,
            {
                "fecha_desde": "2026-07-01",
                "fecha_hasta": "2026-07-31",
                "cliente": otro.id,
                "search": "OC-XLSX",
            },
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(
            'filename="cobranzas_2026-07-01_a_2026-07-31.xlsx"',
            response["Content-Disposition"],
        )
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        sheet = workbook["Cobranzas"]
        headers = [cell.value for cell in sheet[1]]
        self.assertIn("Importe nominal", headers)
        self.assertIn("Importe contable", headers)
        self.assertIn("Factura original asociada", headers)
        self.assertEqual(sheet.max_row, 3)
        filas = {
            sheet.cell(row, 6).value: row
            for row in range(2, sheet.max_row + 1)
        }
        fila_factura = filas[501]
        fila_nota = filas[503]
        self.assertEqual(
            sheet.cell(fila_factura, 1).value.date(),
            date(2026, 7, 31),
        )
        self.assertEqual(
            Decimal(str(sheet.cell(fila_factura, 9).value)),
            Decimal("125430.52"),
        )
        self.assertEqual(
            Decimal(str(sheet.cell(fila_factura, 12).value)),
            Decimal("0.02"),
        )
        self.assertEqual(sheet.cell(fila_nota, 10).value, -1)
        self.assertEqual(
            Decimal(str(sheet.cell(fila_nota, 11).value)),
            Decimal("-0.01"),
        )
        self.assertEqual(
            sheet.cell(fila_nota, 22).value,
            incluida.numero_completo,
        )
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertTrue(sheet.auto_filter.ref)

    def test_exportacion_sin_filtros_permisos_y_licencia(self):
        self.crear_factura(numero_factura=510, total=Decimal("0.01"))
        url = reverse("cobranzas-facturas-exportar-excel")
        response = self.client.get(url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        usuario_sin_permiso = get_user_model().objects.create_user(
            "excel-sin-permiso", is_staff=True
        )
        self.client.force_authenticate(usuario_sin_permiso)
        self.assertEqual(
            self.client.get(url).status_code,
            status.HTTP_403_FORBIDDEN,
        )

        self.client.force_authenticate(self.usuario)
        with patch(
            "licensing.decorators.license_manager.is_enabled",
            return_value=False,
        ):
            sin_licencia = self.client.get(url)
        self.assertEqual(sin_licencia.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(sin_licencia.data["error"], "license_required")
