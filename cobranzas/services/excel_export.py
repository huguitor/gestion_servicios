from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = [
    "Fecha",
    "Cliente",
    "CUIT / identificación",
    "Tipo de comprobante",
    "Punto de venta",
    "Número",
    "Número completo",
    "Orden de compra",
    "Importe nominal",
    "Signo",
    "Importe contable",
    "Total cobrado",
    "Saldo pendiente",
    "Estado",
    "Semáforo",
    "Fecha de envío",
    "Medio de envío",
    "Plazo de cobro",
    "Fecha estimada de cobro",
    "Presupuesto",
    "Remitos",
    "Factura original asociada",
    "Observaciones",
]

MONEY_COLUMNS = (9, 11, 12, 13)
DATE_COLUMNS = (1, 16, 19)
WIDTHS = {
    1: 12,
    2: 30,
    3: 20,
    4: 22,
    7: 18,
    8: 20,
    20: 20,
    21: 28,
    22: 24,
    23: 40,
}


def generar_excel_cobranzas(comprobantes):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cobranzas"
    sheet.append(HEADERS)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center")

    for comprobante in comprobantes:
        presupuesto = comprobante.presupuesto_referencia
        if not presupuesto and comprobante.presupuesto:
            presupuesto = str(comprobante.presupuesto.numero or "")
        remitos = ", ".join(
            remito.numero_formateado for remito in comprobante.remitos.all()
        )
        original = (
            comprobante.comprobante_original.numero_completo
            if comprobante.comprobante_original_id
            else ""
        )
        sheet.append(
            [
                comprobante.fecha_factura,
                str(comprobante.cliente),
                comprobante.cliente.documento or "",
                comprobante.get_tipo_comprobante_display(),
                comprobante.punto_venta,
                comprobante.numero_factura,
                comprobante.numero_completo,
                comprobante.orden_compra,
                comprobante.total,
                comprobante.signo,
                comprobante.importe_con_efecto,
                comprobante.total_cobrado,
                comprobante.saldo_pendiente,
                comprobante.estado,
                comprobante.semaforo,
                comprobante.fecha_envio,
                comprobante.get_medio_envio_display() if comprobante.medio_envio else "",
                comprobante.plazo_cobro_dias,
                comprobante.fecha_estimada_cobro,
                presupuesto,
                remitos,
                original,
                comprobante.observaciones,
            ]
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for column in MONEY_COLUMNS:
            row[column - 1].number_format = '#,##0.00'
        for column in DATE_COLUMNS:
            row[column - 1].number_format = "DD/MM/YYYY"
    for column in range(1, len(HEADERS) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = WIDTHS.get(column, 16)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
